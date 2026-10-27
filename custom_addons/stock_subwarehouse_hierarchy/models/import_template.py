from io import BytesIO
import re

from odoo import api, fields, models
from odoo.exceptions import ValidationError


MRP_COLUMNS = [
    ("name", "制造单号"), ("product_id", "产品"), ("product_qty", "数量"),
    ("product_uom_id", "计量单位"), ("bom_id", "物料清单"), ("origin", "源单据"),
    ("date_start", "计划日期"), ("location_src_id", "组件库位"),
    ("location_dest_id", "成品库位"), ("picking_type_id", "作业类型"),
    ("company_id", "公司"), ("never_product_template_attribute_value_ids", "产品属性值"),
    ("user_id", "负责人"), ("priority", "优先级"),
    ("date_deadline", "截止日期"), ("date_finished", "计划完成日期"),
]


class ImportTemplate(models.Model):
    _name = "stock.subwarehouse.import.template"
    _description = "导入模板配置"

    name = fields.Char(string="模板名称", required=True, default="制造单导入模板")
    res_model = fields.Selection([("mrp.production", "制造单")], required=True, default="mrp.production")
    column_ids = fields.One2many("stock.subwarehouse.import.template.column", "template_id", string="模板列")
    _model_unique = models.Constraint("UNIQUE(res_model)", "每种导入类型只能有一个模板配置。")

    @api.model
    def _manufacturing(self):
        return self.env.ref("stock_subwarehouse_hierarchy.import_template_manufacturing")

    def _columns(self):
        self.ensure_one()
        return self.column_ids.filtered("enabled").sorted(lambda line: (line.sequence, line.id))

    def action_download(self):
        return {"type": "ir.actions.act_url", "url": "/stock_subwarehouse_hierarchy/import_template/mrp_production.xlsx", "target": "self"}

    def _workbook(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        columns = self._columns()
        if not columns:
            raise ValidationError("请管理员至少启用一列模板字段。")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "制造单导入"
        sheet.append(columns.mapped("field_name"))
        sheet.append(columns.mapped("title"))
        sheet.freeze_panes = "A3"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in sheet[2]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E2F0D9")
        for index, column in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = max(18, min(45, len(column.field_name) + 2))
        notes = workbook.create_sheet("填写说明")
        for row in [
            ["项目", "说明"],
            ["标题行", "第1行为字段标识，第2行固定为中文标题。请从第3行填写数据；导入时可另选起始行。"],
            ["制造单号", "删除制造单号列可每行自动编号；保留该列时，可选择将空单号行作为上一单的附加产品。"],
            ["默认值", "缺失或无法识别的数据可生成默认值建议。请核对建议并接受后再导入。"],
            ["关联字段", "填写唯一名称或产品编码；也可把字段标识改为 字段/.id 填数据库ID，或 字段/id 填外部ID。"],
            ["新成品库位", "导入前显示所有新库位及上级仓库；选择全部创建，或跳过受影响行。"],
            ["必填规则", "模板必填由管理员设置。Odoo要求的基础字段缺省时，仍需有效默认值。"],
        ]:
            notes.append(row)
        notes.append(["字段", "中文标题", "模板必填", "默认值建议"])
        for column in columns:
            notes.append([column.field_name, column.title, "是" if column.required else "否", column.default_value or "系统默认"])
        notes.column_dimensions["A"].width = 35
        notes.column_dimensions["B"].width = 100
        notes.column_dimensions["C"].width = 15
        notes.column_dimensions["D"].width = 35
        attributes = workbook.create_sheet("产品属性")
        attributes.append(["属性ID", "属性", "变体创建方式", "值ID", "值"])
        for attribute in self.env["product.attribute"].search([], order="sequence, id"):
            for value in attribute.value_ids:
                attributes.append([attribute.id, attribute.display_name, attribute.create_variant, value.id, value.display_name])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()


class ImportTemplateColumn(models.Model):
    _name = "stock.subwarehouse.import.template.column"
    _description = "导入模板列"
    _order = "sequence, id"

    template_id = fields.Many2one("stock.subwarehouse.import.template", required=True, ondelete="cascade")
    sequence = fields.Integer(default=10)
    enabled = fields.Boolean(string="启用", default=True)
    field_name = fields.Selection(MRP_COLUMNS, string="字段", required=True)
    title = fields.Char(string="中文标题", required=True)
    required = fields.Boolean(string="必填", default=False)
    default_value = fields.Char(string="默认值建议", help="填写与表格相同的值；关联字段填写唯一名称或产品编码。留空使用系统建议。")
    _column_unique = models.Constraint("UNIQUE(template_id, field_name)", "同一模板不能重复添加相同字段。")

    @api.onchange("field_name")
    def _onchange_field_name(self):
        if self.field_name:
            self.title = dict(MRP_COLUMNS)[self.field_name]

    @api.constrains("title")
    def _check_chinese_title(self):
        for column in self:
            if not re.search(r"[\u3400-\u9fff]", column.title or ""):
                raise ValidationError("模板第2行必须使用中文标题，请填写含中文的列标题。")
