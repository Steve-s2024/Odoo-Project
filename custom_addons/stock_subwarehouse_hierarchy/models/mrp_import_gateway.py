import base64
import csv
import hashlib
import json
import math
import re
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO, StringIO

from psycopg2 import IntegrityError

from odoo import api, fields, models, Command
from odoo.exceptions import AccessError, UserError, ValidationError

from .import_template import MRP_COLUMNS


_INTERNAL = object()
_INPUTS = {"file", "filename", "start_row", "sheet_number", "group_blank_names", "company_id", "picking_type_id", "location_id"}
_LABELS = dict(MRP_COLUMNS)


def _text(value):
    if value is None or value is False:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


class _Resolver:
    """One allowed-record query per related model; never create master data."""
    def __init__(self, env):
        self.env = env
        self.cache = {}
        # Metadata is not normally readable by importers. Resolve IDs only;
        # every target is intersected with records visible under their own ACLs.
        self.xmlids = {(r.model, r.complete_name): r.res_id for r in env["ir.model.data"].sudo().search([
            ("model", "in", ["product.product", "uom.uom", "mrp.bom", "stock.location", "stock.picking.type", "res.company", "res.users", "product.template.attribute.value"]),
        ])}

    def records(self, field_name, company_id, product=None, picking=None):
        field = self.env["mrp.production"]._fields[field_name]
        model = self.env[field.comodel_name]
        if field.comodel_name not in self.cache:
            domain = []
            if "company_id" in model._fields:
                domain += [("company_id", "in", [False, *self.env.companies.ids])]
            if model._name == "res.company":
                domain += [("id", "in", self.env.companies.ids)]
            if model._name == "stock.location":
                domain += [("usage", "=", "internal")]
            if model._name == "stock.picking.type":
                domain += [("code", "=", "mrp_operation")]
            if model._name == "product.product":
                domain += [("type", "!=", "service")]
            if model._name == "mrp.bom":
                domain += [("type", "=", "normal")]
            if model._name == "res.users":
                domain += [("share", "=", False)]
            self.cache[model._name] = model.search(domain, order="id")
        records = self.cache[model._name]
        if "company_id" in model._fields:
            records = records.filtered(lambda r: not r.company_id or r.company_id.id == company_id)
        if model._name == "res.users":
            records = records.filtered(lambda r: company_id in r.company_ids.ids)
        if model._name == "mrp.bom" and product:
            records = records.filtered(lambda r: r.product_tmpl_id == product.product_tmpl_id and (not r.product_id or r.product_id == product) and (not r.picking_type_id or r.picking_type_id == picking))
        if model._name == "uom.uom" and product:
            allowed = product.uom_id | product.uom_ids | product.bom_ids.product_uom_id
            records = records & allowed
        if model._name == "product.template.attribute.value" and product:
            records = records.filtered(lambda r: r.product_tmpl_id == product.product_tmpl_id and r.attribute_id.create_variant == "no_variant")
        return records

    def resolve(self, field_name, raw, company_id, suffix="", product=None, picking=None):
        records = self.records(field_name, company_id, product, picking)
        text = _text(raw)
        if suffix == "/.id":
            return records.filtered(lambda r: str(r.id) == text)
        if suffix == "/id":
            record_id = self.xmlids.get((records._name, text if "." in text else "__import__." + text))
            return records.filtered(lambda r: r.id == record_id)
        folded = text.casefold()
        if field_name == "product_id":
            codes = records.filtered(lambda r: (r.default_code or "").casefold() == folded)
            if codes:
                return codes
        return records.filtered(lambda r: folded in {
            _text(r.display_name).casefold(), _text(r.name).casefold() if "name" in r._fields else "",
            _text(r.complete_name).casefold() if "complete_name" in r._fields else "",
        })


class MrpImportGateway(models.TransientModel):
    _name = "stock.subwarehouse.mrp.import.gateway"
    _description = "制造单标准导入"

    file = fields.Binary(string="导入文件", required=True, attachment=False)
    filename = fields.Char(string="文件名")
    start_row = fields.Integer(string="数据起始行（Excel行号）", default=3, required=True)
    sheet_number = fields.Integer(string="工作表序号", default=1, required=True)
    group_blank_names = fields.Boolean(string="空单号行作为上一单的附加产品", default=True)
    company_id = fields.Many2one("res.company", string="默认公司", required=True, default=lambda self: self.env.company)
    picking_type_id = fields.Many2one("stock.picking.type", string="默认制造作业类型", required=True, domain="[('code', '=', 'mrp_operation'), ('company_id', '=', company_id)]")
    location_id = fields.Many2one("stock.location", string="默认制造仓库 / 新库位的上级", required=True, domain="[('usage', '=', 'internal'), ('company_id', 'in', [False, company_id])]")
    state = fields.Selection([("draft", "上传"), ("review", "核对"), ("done", "完成")], default="draft", readonly=True)
    line_ids = fields.One2many("stock.subwarehouse.mrp.import.gateway.line", "gateway_id", string="逐行预览", readonly=False)
    missing_destinations = fields.Text(string="待创建的成品库位", readonly=True)
    destination_decision = fields.Selection([("ask", "请选择处理方式"), ("create", "创建上述全部库位"), ("skip", "跳过使用这些库位的行")], default="ask", string="新库位处理")
    summary = fields.Text(string="检查结果", readonly=True)
    plan_json = fields.Json(readonly=True)
    plan_hash = fields.Char(readonly=True)
    result_batch_id = fields.Many2one("stock.subwarehouse.business.import.result.batch", readonly=True)

    @api.model
    def default_get(self, names):
        values = super().default_get(names)
        production = self.env["mrp.production"].default_get(["picking_type_id", "location_dest_id"])
        if "picking_type_id" in names:
            values.setdefault("picking_type_id", production.get("picking_type_id"))
        if "location_id" in names:
            values.setdefault("location_id", self.env.context.get("subwarehouse_manufacturing_location_id") or self.env.context.get("default_location_dest_id") or production.get("location_dest_id"))
        return values

    @api.model_create_multi
    def create(self, vals_list):
        for vals in vals_list:
            if vals.get("line_ids") == []:
                vals.pop("line_ids")
            if vals.get("state") == "draft":
                vals.pop("state")
            if vals.get("destination_decision") == "ask":
                vals.pop("destination_decision")
            if set(vals) - _INPUTS:
                raise AccessError("只能设置导入文件及导入选项。")
        return super().create(vals_list)

    def write(self, vals):
        if set(vals) - (_INPUTS | {"destination_decision", "line_ids"}):
            raise AccessError("导入预览由服务器生成，不能直接修改。")
        if any(w.state == "done" for w in self):
            raise UserError("此文件已处理，请新建导入任务。")
        # The UI may only change existing rows' inclusion, never insert a plan.
        if "line_ids" in vals:
            for command in vals["line_ids"]:
                if command[0] != Command.UPDATE or set(command[2]) != {"include"}:
                    raise AccessError("只能选择是否导入已有预览行。")
                if command[1] not in self.line_ids.ids:
                    raise AccessError("预览行不属于此导入任务。")
        result = super().write(vals)
        if set(vals) & _INPUTS:
            super().write({"state": "draft", "plan_hash": False})
        return result

    def _save(self, vals):
        return super().write(vals)

    def _lock(self):
        self.ensure_one()
        self.check_access("write")
        if not self.env.su and self.create_uid != self.env.user:
            raise AccessError("只能处理自己上传的导入文件。")
        self.env.cr.execute("SELECT id FROM stock_subwarehouse_mrp_import_gateway WHERE id = %s FOR UPDATE", [self.id])
        self.invalidate_recordset()

    def _context_model(self):
        if self.company_id not in self.env.companies:
            raise AccessError("默认公司不在当前允许访问的公司范围内。")
        if self.picking_type_id.company_id != self.company_id or self.picking_type_id.code != "mrp_operation":
            raise ValidationError("请选择默认公司的制造作业类型。")
        if not self.location_id.active or self.location_id.usage != "internal" or (self.location_id.company_id and self.location_id.company_id != self.company_id):
            raise ValidationError("请选择默认公司的有效内部仓库。")
        return self.env["mrp.production"].with_company(self.company_id).with_context(
            default_picking_type_id=self.picking_type_id.id,
            default_location_dest_id=self.location_id.id, default_location_src_id=self.location_id.id,
            # Per-row company/location choices must take precedence over the launch context.
            subwarehouse_manufacturing_location_id=False,
        )

    def _fingerprint(self):
        template = self.env["stock.subwarehouse.import.template"]._manufacturing()
        values = [self.file.decode() if isinstance(self.file, bytes) else self.file,
                  self.filename, self.start_row, self.sheet_number, self.group_blank_names,
                  self.company_id.id, self.picking_type_id.id, self.location_id.id,
                  self.env.uid, self.env.companies.ids,
                  [(c.field_name, c.title, c.required, c.default_value) for c in template._columns()]]
        return hashlib.sha256(json.dumps(values, ensure_ascii=False).encode()).hexdigest()

    def _read_rows(self):
        if self.start_row < 2 or self.sheet_number < 1:
            raise ValidationError("第1行是字段标识，数据起始行至少为2；工作表序号至少为1。")
        content = base64.b64decode(self.file or b"")
        if not content or len(content) > 20 * 1024 * 1024:
            raise ValidationError("请上传不超过20MB的 XLSX 或 CSV 文件。")
        if (self.filename or "").lower().endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                if self.sheet_number > len(workbook.worksheets):
                    raise ValidationError("所选工作表不存在。")
                iterator = workbook.worksheets[self.sheet_number - 1].iter_rows(values_only=True)
                rows = []
                for index, row in enumerate(iterator, 1):
                    if index > 10002:
                        raise ValidationError("每次最多读取10000行数据，请拆分文件。")
                    rows.append([_text(v) for v in row])
            finally:
                workbook.close()
        elif (self.filename or "").lower().endswith(".csv"):
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("gb18030")
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(StringIO(text), dialect))
            if len(rows) > 10002:
                raise ValidationError("每次最多读取10000行数据，请拆分文件。")
        else:
            raise ValidationError("请选择 XLSX 或 CSV 文件。")
        if not rows:
            raise ValidationError("文件为空。")
        columns = self.env["stock.subwarehouse.import.template"]._manufacturing()._columns()
        aliases = {c.title: c.field_name for c in columns}
        enabled = set(columns.mapped("field_name"))
        headers, seen = [], set()
        for value in rows[0]:
            header = aliases.get(_text(value), _text(value))
            base, sep, suffix = header.partition("/")
            if not header:
                headers.append(None)
                continue
            if base not in enabled or (sep and suffix not in {"id", ".id"}):
                raise ValidationError(f"列“{header}”未启用或不受支持。请管理员在导入模板中添加/启用该列，或移除文件中的此列。")
            field = self.env["mrp.production"]._fields[base]
            if sep and field.type != "many2one":
                raise ValidationError(f"列“{header}”不支持关联ID。")
            if base in seen:
                raise ValidationError(f"字段“{base}”重复，请仅保留一列。")
            seen.add(base)
            headers.append(header)
        if not seen:
            raise ValidationError("第1行没有可识别的字段标识。")
        result = []
        for source_row, row in enumerate(rows[self.start_row - 1:], self.start_row):
            if not any(_text(v) for v in row):
                continue
            if any(_text(v) for v in row[len(headers):]) or any(_text(v) for i, v in enumerate(row[:len(headers)]) if not headers[i]):
                raise ValidationError(f"第{source_row}行存在没有字段标识的数据列，请补充第1行的字段标识。")
            result.append({"source_row": source_row, "raw": {h: _text(row[i]) if i < len(row) else "" for i, h in enumerate(headers) if h}})
        if not result:
            raise ValidationError("所选起始行之后没有数据。")
        return result, seen

    def action_download_template(self):
        return self.env["stock.subwarehouse.import.template"]._manufacturing().action_download()

    def action_configure_template(self):
        if not self.env.user.has_group("base.group_system"):
            raise AccessError("仅管理员可配置导入模板。")
        return {"type": "ir.actions.act_window", "res_model": "stock.subwarehouse.import.template", "res_id": self.env["stock.subwarehouse.import.template"]._manufacturing().id, "view_mode": "form", "target": "current"}

    def action_preview(self):
        return self._preview(False)

    def action_generate_defaults(self):
        return self._preview(True)

    def _normalize(self, item, resolver, config, generate, parent=None):
        raw = item["raw"]
        raw_by_field = {key.split("/")[0]: (value, key[len(key.split("/")[0]):]) for key, value in raw.items()}
        values, changes, errors = {}, [], []
        item.update(values=values, changes=changes, errors=errors)
        company_id = parent["values"].get("company_id", self.company_id.id) if parent else self.company_id.id
        product, picking = None, self.picking_type_id
        order = ["company_id", "picking_type_id", "product_id", "product_qty", "product_uom_id", "location_src_id", "location_dest_id", "bom_id", "name", "origin", "date_start", "date_finished", "date_deadline", "user_id", "priority", "never_product_template_attribute_value_ids"]
        essential = {"company_id", "picking_type_id", "product_id", "product_qty", "product_uom_id", "location_src_id", "location_dest_id", "date_start"}
        for key in order:
            if parent and key not in {"product_id", "product_qty", "product_uom_id", "location_dest_id"}:
                continue
            original, suffix = raw_by_field.get(key, ("", ""))
            column = config.get(key)
            required = key in essential or bool(column and column.required)
            if not original and not required:
                continue
            title = column.title if column else _LABELS[key]
            field = self.env["mrp.production"]._fields[key]
            # Named destinations are collected across the entire file below.
            if key == "location_dest_id" and original:
                item["destination"] = (original, suffix)
                continue

            def parse(value, value_suffix=""):
                if field.type == "many2one":
                    match = resolver.resolve(key, value, company_id, value_suffix, product, picking)
                    if len(match) != 1:
                        raise ValueError("不存在、重复或不适用于该制造单")
                    return match.id, match.display_name
                if field.type in {"float", "integer"}:
                    text = _text(value)
                    if "," in text and not re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?", text):
                        raise ValueError("数字格式无法识别，请使用小数点或规范的千分位")
                    number = float(text.replace(",", ""))
                    if not math.isfinite(number) or (key == "product_qty" and number <= 0):
                        raise ValueError("数量必须为正数")
                    return number, str(number)
                if field.type in {"date", "datetime"}:
                    parsed = datetime.fromisoformat(_text(value).replace("/", "-"))
                    if parsed.tzinfo:
                        raise ValueError("请使用无时区的本地时间")
                    if field.type == "datetime":
                        import pytz
                        zone = pytz.timezone(self.env.user.tz or "UTC")
                        parsed = zone.localize(parsed).astimezone(pytz.UTC).replace(tzinfo=None)
                    return fields.Datetime.to_string(parsed), _text(value)
                if field.type == "selection":
                    choices = dict(field._description_selection(self.env))
                    value = next((k for k, label in choices.items() if _text(value) in {k, label}), None)
                    if value is None:
                        raise ValueError("不是有效选项")
                    return value, choices[value]
                if field.type == "many2many":
                    records = self.env[field.comodel_name]
                    for token in _text(value).split(","):
                        match = resolver.resolve(key, token, company_id, product=product)
                        if len(match) != 1:
                            raise ValueError("属性值不存在、重复或不适用于该产品")
                        records |= match
                    return [Command.set(records.ids)], ", ".join(records.mapped("display_name"))
                if not _text(value):
                    raise ValueError("缺少数据")
                return _text(value), _text(value)

            reason = "缺少数据"
            try:
                if not original:
                    raise ValueError(reason)
                value, display = parse(original, suffix)
            except (ValueError, TypeError, OverflowError) as error:
                reason = str(error)
                if not generate:
                    errors.append(f"{title}：{reason}。可生成默认值建议，或修改文件。")
                    continue
                fallback = column.default_value if column else False
                fallback_suffix = ""
                if not fallback:
                    if field.type == "many2one":
                        candidates = resolver.records(key, company_id, product, picking)
                        preferred = {
                            "company_id": company_id, "picking_type_id": self.picking_type_id.id,
                            "product_uom_id": product.uom_id.id if product else False,
                            "location_src_id": self.location_id.id if company_id == self.company_id.id else picking.default_location_src_id.id,
                            "location_dest_id": (parent["values"].get("location_dest_id") if parent else False) or (self.location_id.id if company_id == self.company_id.id else picking.default_location_dest_id.id),
                            "user_id": self.env.uid,
                        }.get(key)
                        selected = candidates.filtered(lambda r: r.id == preferred)[:1] or candidates[:1]
                        fallback, fallback_suffix = str(selected.id) if selected else "", "/.id"
                    elif key == "product_qty":
                        fallback = "1"
                    elif field.type == "selection":
                        fallback = field._description_selection(self.env)[0][0]
                    elif field.type in {"date", "datetime"}:
                        fallback = fields.Datetime.context_timestamp(self, fields.Datetime.now()).replace(tzinfo=None).strftime("%Y-%m-%d %H:%M:%S")
                    elif key == "name":
                        fallback = ""
                    elif field.type == "many2many":
                        fallback = ""
                    else:
                        fallback = "自动填充"
                try:
                    if not fallback and key == "name":
                        value, display = False, "系统自动编号"
                    elif not fallback and not required:
                        value, display = False, "留空（使用系统默认）"
                    else:
                        value, display = parse(fallback, fallback_suffix)
                except (ValueError, TypeError, OverflowError):
                    errors.append(f"{title}：没有可用的有效默认值，请修改文件或模板默认值。")
                    continue
                changes.append(f"{title}：{original or '（空）'} → {display}（{reason}；默认值建议）")
            if value is not False:
                values[key] = value
            if key == "company_id":
                company_id = value
            elif key == "picking_type_id":
                picking = self.env["stock.picking.type"].browse(value)
            elif key == "product_id":
                product = self.env["product.product"].browse(value)
        item["company_id"] = company_id
        item["picking_type_id"] = picking.id

    def _collect_destinations(self, items, model, generate):
        missing = {}
        for suffix in ("", "/.id", "/id"):
            selected = [r for r in items if r.get("destination", (None, None))[1] == suffix]
            if not selected:
                continue
            data = [[r["destination"][0], str(r["company_id"]), str(r["picking_type_id"])] for r in selected]
            plan = model._inspect_mrp_import_destinations(["location_dest_id" + suffix, "company_id/.id", "picking_type_id/.id"], data)
            for entry in plan["missing"]:
                indexes = entry.pop("source_indexes")
                target = missing.setdefault(entry["key"], {**entry, "rows": []})
                for index in indexes:
                    selected[index]["destination_key"] = entry["key"]
                    target["rows"].append(selected[index]["source_row"])
            for index, item in enumerate(selected):
                if index in plan["resolved"]:
                    item["values"]["location_dest_id"] = plan["resolved"][index]
                elif not item.get("destination_key"):
                    if generate:
                        fallback = self.location_id if item["company_id"] == self.company_id.id else self.env["stock.picking.type"].browse(item["picking_type_id"]).default_location_dest_id
                        if fallback and fallback.active and fallback.usage == "internal":
                            item["values"]["location_dest_id"] = fallback.id
                            item["changes"].append(f"成品库位：{item['destination'][0]} → {fallback.complete_name}（{plan['invalid'].get(index, '无法识别')}；默认值建议）")
                            continue
                    item["errors"].append(plan["invalid"].get(index, "成品库位无法识别。"))
        return list(missing.values())

    def _group_values(self, rows, destination_ids, preview=False):
        values = dict(rows[0]["values"])
        for item in rows:
            if item.get("destination_key"):
                destination_id = destination_ids.get(item["destination_key"])
                if not destination_id:
                    raise ValidationError("待创建的成品库位尚未获准创建。")
                if item == rows[0]:
                    values["location_dest_id"] = destination_id
        if preview and not values.get("name"):
            values["name"] = f"__IMPORT_PREVIEW_{self.id}_{rows[0]['source_row']}__"
        moves = []
        for row in rows[1:]:
            vals = row["values"]
            product = self.env["product.product"].browse(vals["product_id"])
            moves.append(Command.create({
                "product_id": product.id,
                "product_uom_qty": vals["product_qty"], "product_uom": vals["product_uom_id"],
                "location_id": product.with_company(values["company_id"]).property_stock_production.id,
                "location_dest_id": destination_ids.get(row.get("destination_key")) or vals.get("location_dest_id") or values["location_dest_id"],
                "company_id": values["company_id"],
            }))
        if moves:
            values["move_byproduct_ids"] = moves
        return values

    def _existing_names(self):
        self.env.cr.execute("SELECT lower(btrim(name)) FROM mrp_production WHERE name IS NOT NULL")
        return {r[0].casefold() for r in self.env.cr.fetchall()}

    def _validate_current_values(self, rows, destination_ids, resolver):
        company_id = rows[0]["values"]["company_id"]
        if company_id not in self.env.companies.ids:
            raise AccessError("该制造单的公司不在当前允许访问的范围内。")
        picking = self.env["stock.picking.type"].browse(rows[0]["values"]["picking_type_id"])
        for row in rows:
            values = row["values"]
            product = self.env["product.product"].browse(values["product_id"])
            for key, value in values.items():
                field = self.env["mrp.production"]._fields[key]
                if field.type == "many2one" and value not in resolver.records(key, company_id, product, picking).ids:
                    raise ValidationError(f"{_LABELS[key]}已停用、不可访问或不再适用于该制造单，请重新检查。")
            if row.get("destination_key"):
                location = self.env["stock.location"].browse(destination_ids[row["destination_key"]])
                location.check_access("read")
                if not location.active or location.usage != "internal" or (location.company_id and location.company_id.id != company_id):
                    raise ValidationError("成品库位已改变，请重新检查。")

    def _preview(self, generate):
        self._lock()
        if self.state == "done":
            return self.result_batch_id.action_open_lines()
        model = self._context_model()
        model.check_access("create")
        items, supplied = self._read_rows()
        config = {c.field_name: c for c in self.env["stock.subwarehouse.import.template"]._manufacturing()._columns()}
        resolver = _Resolver(self.env)
        groups, parent = [], None
        for item in items:
            is_parent = parent is None or "name" not in supplied or not self.group_blank_names or bool(item["raw"].get("name"))
            if is_parent:
                groups.append([])
                parent = item
            item["group"] = parent["source_row"]
            self._normalize(item, resolver, config, generate, None if is_parent else parent)
            groups[-1].append(item)
        missing = self._collect_destinations(items, model, generate)
        existing, accepted = self._existing_names(), set()
        # Roll back both provisional warehouses and manufacturing records.
        # Use temporary explicit names so previews never consume MO sequences.
        savepoint = self.env.cr.savepoint()
        try:
            destination_ids = {}
            if missing:
                try:
                    with self.env.cr.savepoint():
                        destination_ids = model._create_mrp_import_destinations(missing)
                except (UserError, IntegrityError) as error:
                    for item in items:
                        if item.get("destination_key"):
                            item["errors"].append(f"无法创建成品库位：{error}")
            for group in groups:
                head = group[0]
                name = _text(head["values"].get("name")).casefold()
                if name and (name in existing or name in accepted):
                    head["errors"].append("制造单号已存在于ERP或本次文件中。")
                if name:
                    accepted.add(name)
                if head["errors"]:
                    continue
                valid_rows = [row for row in group if not row["errors"]]
                try:
                    with self.env.cr.savepoint():
                        record = model.with_context(shop_api_skip_event=True, tracking_disable=True).create(self._group_values(valid_rows, destination_ids, preview=True))
                        record.flush_recordset()
                except (UserError, IntegrityError, ValueError) as error:
                    head["errors"].append(f"Odoo校验未通过：{error}")
        finally:
            savepoint.close(rollback=True)
        line_values = []
        for group in groups:
            for item in group:
                if item != group[0] and group[0]["errors"]:
                    item["errors"].append("所属制造单未通过校验，该附加产品行将随制造单跳过。")
                line_values.append({
                    "gateway_id": self.id, "source_row": item["source_row"], "group_row": item["group"],
                    "include": not bool(item["errors"]),
                    "original": "\n".join(f"{_LABELS.get(k.split('/')[0], k)}：{v}" for k, v in item["raw"].items()),
                    "changes": "\n".join(item["changes"]) or "无值替换；缺省制造单号使用系统自动编号。",
                    "error": "\n".join(item["errors"]),
                    "status": "invalid" if item["errors"] else ("proposed" if item["changes"] or item.get("destination_key") else "ready"),
                })
        self.line_ids.with_context(_gateway_token=_INTERNAL).unlink()
        self.env["stock.subwarehouse.mrp.import.gateway.line"].with_context(_gateway_token=_INTERNAL).create(line_values)
        self._save({
            "state": "review", "plan_json": {"items": items, "missing": missing},
            "plan_hash": self._fingerprint(), "destination_decision": "ask",
            "missing_destinations": "\n".join(f"{m['name']} → 上级：{m['parent']}；公司：{m['company']}；Excel行：{', '.join(map(str, m['rows']))}" for m in missing),
            "summary": f"共{len(items)}行，{len(groups)}张制造单；{sum(bool(r['errors']) for r in items)}行未通过校验。核对建议、选择要导入的行后接受并导入。默认值只保证有效，不代表真实业务数据。",
        })
        return self._action()

    def _action(self):
        return {"type": "ir.actions.act_window", "name": "制造单标准导入", "res_model": self._name, "res_id": self.id, "view_mode": "form", "target": "current"}

    def action_accept_import(self):
        self._lock()
        if self.state == "done":
            return self.result_batch_id.action_open_lines()
        if self.state != "review" or self.plan_hash != self._fingerprint():
            raise UserError("文件、选项或模板已改变，请重新检查并核对建议。")
        model = self._context_model()
        model.check_access("create")
        plan = self.plan_json
        if plan["missing"] and self.destination_decision not in {"create", "skip"}:
            raise UserError("请先选择：创建全部新库位，或跳过使用新库位的行。")
        self.env.cr.execute("SELECT pg_advisory_xact_lock(hashtext(%s))", ["stock_subwarehouse_hierarchy.mrp_name_import"])
        existing = self._existing_names()
        destination_ids = model._create_mrp_import_destinations(plan["missing"]) if plan["missing"] and self.destination_decision == "create" else {}
        resolver = _Resolver(self.env)
        batch = self.env["stock.subwarehouse.business.import.result.batch"].create({"import_job_id": -self.id, "res_model": "mrp.production"})
        selected = {line.source_row for line in self.line_ids if line.include}
        groups = defaultdict(list)
        for row in plan["items"]:
            groups[row["group"]].append(row)
        result_lines = []
        for group in groups.values():
            kept, rejected = [], {}
            for row in group:
                reason = "\n".join(row["errors"])
                if row["source_row"] not in selected:
                    reason = reason or "导入人取消勾选此行。"
                if row.get("destination_key") and self.destination_decision == "skip":
                    reason = "按选择跳过使用不存在成品库位的行。"
                if reason:
                    rejected[row["source_row"]] = reason
                else:
                    kept.append(row)
            head = group[0]
            record = model.browse()
            name = _text(head["values"].get("name"))
            if head not in kept:
                for row in kept:
                    rejected[row["source_row"]] = "所属制造单未导入，该附加产品行已跳过。"
                kept = []
            if kept:
                try:
                    with self.env.cr.savepoint():
                        if name and name.casefold() in existing:
                            raise ValidationError("制造单号已存在于ERP或本次导入中。")
                        self._validate_current_values(kept, destination_ids, resolver)
                        record = model.create(self._group_values(kept, destination_ids))
                        record.flush_recordset()
                        existing.add(record.name.strip().casefold())
                except (UserError, IntegrityError, ValueError) as error:
                    record = model.browse()
                    for row in kept:
                        rejected[row["source_row"]] = f"导入时校验未通过：{error}"
            for row in group:
                reason = rejected.get(row["source_row"])
                result_lines.append({"batch_id": batch.id, "source_row": row["source_row"],
                    "status": "failed" if reason else "success", "identifier": record.name if record else name,
                    "record_ref": f"mrp.production,{record.id}" if record and not reason else False,
                    "reason": reason or ("导入成功。" + ("已接受默认值建议：" + "；".join(row["changes"]) if row["changes"] else "")),
                })
        self.env["stock.subwarehouse.business.import.result.line"].create(result_lines)
        self._save({"state": "done", "result_batch_id": batch.id})
        return batch.action_open_lines()


class MrpImportGatewayLine(models.TransientModel):
    _name = "stock.subwarehouse.mrp.import.gateway.line"
    _description = "制造单导入预览行"
    _order = "source_row, id"

    gateway_id = fields.Many2one("stock.subwarehouse.mrp.import.gateway", required=True, ondelete="cascade")
    source_row = fields.Integer(string="Excel行号", readonly=True)
    group_row = fields.Integer(string="所属制造单起始行", readonly=True)
    include = fields.Boolean(string="导入此行", default=True)
    original = fields.Text(string="原始数据", readonly=True)
    changes = fields.Text(string="建议修改（接受后生效）", readonly=True)
    error = fields.Text(string="未解决的问题", readonly=True)
    status = fields.Selection([("ready", "可导入"), ("proposed", "待接受建议"), ("invalid", "需修正或跳过")], string="状态", readonly=True)

    @api.model_create_multi
    def create(self, vals_list):
        if self.env.context.get("_gateway_token") is not _INTERNAL:
            raise AccessError("预览行只能由导入检查生成。")
        return super().create(vals_list)

    def write(self, vals):
        if set(vals) != {"include"}:
            raise AccessError("只能选择是否导入此行。")
        if any(line.gateway_id.state != "review" for line in self):
            raise UserError("请先生成预览；已完成的导入不能更改。")
        return super().write(vals)

    def unlink(self):
        if self.env.context.get("_gateway_token") is not _INTERNAL and not self.env.su:
            raise AccessError("预览行只能通过重新检查更新。")
        return super().unlink()


class MrpProductionGateway(models.Model):
    _inherit = "mrp.production"

    @api.model
    def _get_dynamic_import_template_columns(self):
        return [(c.field_name, c.title) for c in self.env["stock.subwarehouse.import.template"]._manufacturing()._columns()]

    @api.model
    def _generate_dynamic_import_template_xlsx(self):
        return self.env["stock.subwarehouse.import.template"]._manufacturing()._workbook()


class StockLocationGateway(models.Model):
    _inherit = "stock.location"

    def action_import_manufacturing_sheet(self):
        action = super().action_import_manufacturing_sheet()
        return {"type": "ir.actions.act_window", "name": "制造单标准导入", "res_model": "stock.subwarehouse.mrp.import.gateway", "view_mode": "form", "target": "current", "context": action["context"]}
