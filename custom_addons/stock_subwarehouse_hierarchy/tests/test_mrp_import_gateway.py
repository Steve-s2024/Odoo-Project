import base64
import csv
from io import BytesIO, StringIO

from odoo import Command
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.tests import TransactionCase, tagged, Form


@tagged("post_install", "-at_install", "mrp_import_gateway")
class TestMrpImportGateway(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.env = cls.env(context=dict(cls.env.context, shop_api_skip_event=True, tracking_disable=True))
        cls.warehouse = cls.env["stock.warehouse"].create({"name": "Gateway Test", "code": "GWT"})
        cls.product = cls.env["product.product"].create({"name": "Gateway Product", "default_code": "GW-PRODUCT", "is_storable": True})
        cls.other = cls.env["product.product"].create({"name": "Gateway Other", "default_code": "GW-OTHER", "is_storable": True})
        cls.template = cls.env["stock.subwarehouse.import.template"]._manufacturing()

    def gateway(self, data, headers=None, start_row=3, **kwargs):
        headers = headers or ["product_id", "product_qty", "location_dest_id"]
        text = StringIO()
        writer = csv.writer(text)
        writer.writerow(headers)
        writer.writerow(["中文标题"] * len(headers))
        writer.writerows(data)
        return self.env["stock.subwarehouse.mrp.import.gateway"].create({
            "file": base64.b64encode(text.getvalue().encode()), "filename": "manufacturing.csv",
            "start_row": start_row, "company_id": self.env.company.id,
            "picking_type_id": self.warehouse.manu_type_id.id,
            "location_id": self.warehouse.lot_stock_id.id, **kwargs,
        })

    def results(self, wizard):
        return wizard.result_batch_id.line_ids

    def test_defaults_are_reviewed_preview_rolls_back_and_confirm_is_idempotent(self):
        wizard = self.gateway([["GW-PRODUCT", "bad", ""]])
        before = self.env["mrp.production"].search_count([])
        wizard.action_preview()
        self.assertEqual(wizard.line_ids.status, "invalid")
        wizard.action_generate_defaults()
        self.assertFalse(wizard.line_ids.error, wizard.line_ids.error)
        self.assertIn("bad → 1.0", wizard.line_ids.changes)
        self.assertEqual(self.env["mrp.production"].search_count([]), before)
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).status, "success", self.results(wizard).reason)
        order = self.results(wizard).record_ref
        self.assertEqual(order.product_qty, 1)
        self.assertNotIn("__IMPORT_PREVIEW", order.name)
        wizard.action_accept_import()
        self.assertEqual(self.env["mrp.production"].search_count([]), before + 1)

    def test_configure_columns_titles_required_and_default(self):
        self.template.column_ids.filtered(lambda c: c.field_name == "name").unlink()
        column = self.template.column_ids.filtered(lambda c: c.field_name == "origin")
        column.write({"required": True, "default_value": "导入默认来源", "title": "来源说明", "sequence": 1})
        added = self.env["stock.subwarehouse.import.template.column"].create({"template_id": self.template.id, "field_name": "priority", "title": "处理优先级", "sequence": 2})
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(self.env["mrp.production"]._generate_dynamic_import_template_xlsx()))
        rows = list(workbook["制造单导入"].values)
        self.assertEqual(rows[0][:2], ("origin", "priority"))
        self.assertEqual(rows[1][:2], ("来源说明", "处理优先级"))
        self.assertNotIn("name", rows[0])
        wizard = self.gateway([["GW-PRODUCT", "2", ""]])
        wizard.action_generate_defaults()
        self.assertIn("导入默认来源", wizard.line_ids.changes)
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).record_ref.origin, "导入默认来源")
        added.enabled = False
        self.assertNotIn("priority", [c[0] for c in self.env["mrp.production"]._get_dynamic_import_template_columns()])
        with self.assertRaises(ValidationError), self.env.cr.savepoint():
            column.title = "English only"

    def test_optional_column_omitted_uses_normal_numbering(self):
        wizard = self.gateway([["GW-PRODUCT", "1", ""], ["GW-PRODUCT", "2", ""]])
        wizard.action_generate_defaults()
        wizard.action_accept_import()
        self.assertEqual(set(self.results(wizard).mapped("status")), {"success"})
        self.assertEqual(len(set(self.results(wizard).mapped("identifier"))), 2)

    def test_configured_product_default_replaces_unrecognized_product(self):
        self.template.column_ids.filtered(lambda c: c.field_name == "product_id").default_value = "GW-PRODUCT"
        wizard = self.gateway([["unknown-product", "1", ""]])
        wizard.action_generate_defaults()
        self.assertIn("unknown-product", wizard.line_ids.changes)
        self.assertIn("Gateway Product", wizard.line_ids.changes)
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).record_ref.product_id, self.product)

    def test_start_row_and_blank_lines_keep_physical_numbers(self):
        wizard = self.gateway([["instructions", "", ""], ["GW-PRODUCT", "1", ""], ["", "", ""], ["GW-PRODUCT", "2", ""]], start_row=4)
        wizard.action_generate_defaults()
        self.assertEqual(wizard.line_ids.mapped("source_row"), [4, 6])
        wizard.line_ids[-1].include = False
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).mapped("source_row"), [4, 6])
        self.assertEqual(self.results(wizard).mapped("status"), ["success", "failed"])

    def test_missing_destinations_create_all_only_after_acceptance(self):
        wizard = self.gateway([["GW-PRODUCT", "1", "Gateway A"], ["GW-PRODUCT", "1", "Gateway B"]])
        before = self.env["stock.location"].search_count([])
        wizard.action_generate_defaults()
        self.assertEqual(len(wizard.plan_json["missing"]), 2)
        self.assertEqual(self.env["stock.location"].search_count([]), before)
        with self.assertRaises(UserError):
            wizard.action_accept_import()
        wizard.destination_decision = "create"
        wizard.action_accept_import()
        self.assertEqual(self.env["stock.location"].search_count([]), before + 2)
        self.assertEqual(self.results(wizard).mapped("status"), ["success", "success"], self.results(wizard).mapped("reason"))

    def test_skip_unknown_destination_preserves_valid_rows(self):
        wizard = self.gateway([["GW-PRODUCT", "1", "Gateway skipped"], ["GW-PRODUCT", "1", self.warehouse.lot_stock_id.complete_name]])
        wizard.action_generate_defaults()
        wizard.destination_decision = "skip"
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).mapped("status"), ["failed", "success"])
        self.assertFalse(self.env["stock.location"].search([("name", "=", "Gateway skipped")]))

    def test_duplicate_orders_partial_success(self):
        wizard = self.gateway([["GW-UNIQUE", "GW-PRODUCT", "1"], [" gw-unique ", "GW-PRODUCT", "1"], ["GW-SECOND", "GW-PRODUCT", "1"]], headers=["name", "product_id", "product_qty"])
        wizard.action_generate_defaults()
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).mapped("status"), ["success", "failed", "success"])

    def test_grouped_products_are_one_order_and_invalid_child_skips(self):
        wizard = self.gateway([["GW-GROUP", "GW-PRODUCT", "1"], ["", "GW-OTHER", "2"]], headers=["name", "product_id", "product_qty"])
        wizard.action_generate_defaults()
        self.assertFalse(any(wizard.line_ids.mapped("error")), wizard.line_ids.mapped("error"))
        wizard.action_accept_import()
        self.assertEqual(set(self.results(wizard).mapped("status")), {"success"}, self.results(wizard).mapped("reason"))
        self.assertEqual(len(set(self.results(wizard).mapped("identifier"))), 1)
        self.assertEqual(self.results(wizard)[0].record_ref.move_byproduct_ids.product_id, self.other)

    def test_file_or_template_change_requires_new_preview(self):
        wizard = self.gateway([["GW-PRODUCT", "1", ""]])
        wizard.action_generate_defaults()
        self.template.column_ids[0].required = not self.template.column_ids[0].required
        with self.assertRaises(UserError):
            wizard.action_accept_import()
        wizard.action_generate_defaults()
        wizard.start_row = 4
        with self.assertRaises(UserError):
            wizard.action_accept_import()

    def test_plan_cannot_be_forged(self):
        wizard = self.gateway([["GW-PRODUCT", "1", ""]])
        wizard.action_generate_defaults()
        with self.assertRaises(AccessError):
            wizard.write({"plan_json": {"items": []}})
        with self.assertRaises(AccessError):
            wizard.line_ids.write({"error": False})
        with self.assertRaises(AccessError):
            self.env["stock.subwarehouse.mrp.import.gateway.line"].create({"gateway_id": wizard.id})

    def test_changed_master_data_is_revalidated(self):
        wizard = self.gateway([["GW-PRODUCT", "1", ""]])
        wizard.action_generate_defaults()
        self.product.active = False
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).status, "failed")
        self.assertIn("停用", self.results(wizard).reason)

    def test_ordinary_importer_cannot_edit_admin_template(self):
        user = self.env["res.users"].create({"name": "Gateway Importer", "login": "gateway-importer", "group_ids": [Command.set([self.env.ref("mrp.group_mrp_user").id, self.env.ref("stock.group_stock_user").id])]})
        with self.assertRaises(AccessError):
            self.template.with_user(user).write({"name": "forbidden"})
        with self.assertRaises(AccessError):
            self.template.column_ids[0].with_user(user).write({"required": False})
        source = self.gateway([["GW-PRODUCT", "1", ""]])
        wizard = self.env["stock.subwarehouse.mrp.import.gateway"].with_user(user).create({
            "file": source.file, "filename": source.filename, "company_id": self.env.company.id,
            "picking_type_id": self.warehouse.manu_type_id.id, "location_id": self.warehouse.lot_stock_id.id,
        })
        wizard.action_generate_defaults()
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).status, "success", self.results(wizard).reason)
        other_user = self.env["res.users"].create({"name": "Gateway Other User", "login": "gateway-other-user", "group_ids": [Command.set(user.group_ids.ids)]})
        with self.assertRaises(AccessError):
            wizard.with_user(other_user).action_generate_defaults()

    def test_xlsx_title_row_and_selected_sheet(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.title = "说明"
        sheet = workbook.create_sheet("导入")
        sheet.append(["product_id", "product_qty"])
        sheet.append(["产品", "数量"])
        sheet.append(["GW-PRODUCT", 2])
        output = BytesIO()
        workbook.save(output)
        wizard = self.gateway([["placeholder", "", ""]])
        wizard.write({"file": base64.b64encode(output.getvalue()), "filename": "test.xlsx", "sheet_number": 2})
        wizard.action_generate_defaults()
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).status, "success", self.results(wizard).reason)
        self.assertEqual(self.results(wizard).record_ref.product_qty, 2)

    def test_form_upload_and_defaults(self):
        source = self.gateway([["GW-PRODUCT", "1", ""]])
        with Form(self.env["stock.subwarehouse.mrp.import.gateway"].with_context(
            default_picking_type_id=self.warehouse.manu_type_id.id,
            default_location_dest_id=self.warehouse.lot_stock_id.id,
            default_filename="form.csv",
        )) as form:
            form.file = source.file
        wizard = form.record
        wizard.action_generate_defaults()
        self.assertFalse(wizard.line_ids.error, wizard.line_ids.error)

    def test_skip_invalid_continuation_preserves_main_order(self):
        wizard = self.gateway([["GW-SKIP-CHILD", "GW-PRODUCT", "1", ""], ["", "GW-OTHER", "1", "Missing child"]], headers=["name", "product_id", "product_qty", "location_dest_id"])
        wizard.action_generate_defaults()
        wizard.destination_decision = "skip"
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).mapped("status"), ["success", "failed"])
        self.assertFalse(self.results(wizard)[0].record_ref.move_byproduct_ids)

    def test_no_available_configured_default_skips_row(self):
        self.template.column_ids.filtered(lambda c: c.field_name == "product_id").default_value = "NONEXISTENT DEFAULT"
        wizard = self.gateway([["unknown", "1", ""], ["GW-PRODUCT", "1", ""]])
        wizard.action_generate_defaults()
        wizard.action_accept_import()
        self.assertEqual(self.results(wizard).mapped("status"), ["failed", "success"])

    def test_source_file_database_and_external_ids(self):
        xmlid = "stock_subwarehouse_hierarchy.gateway_test_product"
        self.env["ir.model.data"].create({"module": xmlid.split(".")[0], "name": xmlid.split(".")[1], "model": "product.product", "res_id": self.product.id})
        for header, value in [("product_id/.id", str(self.product.id)), ("product_id/id", xmlid)]:
            wizard = self.gateway([[value, "1"]], headers=[header, "product_qty"])
            wizard.action_generate_defaults()
            wizard.action_accept_import()
            self.assertEqual(self.results(wizard).status, "success", self.results(wizard).reason)
            self.assertEqual(self.results(wizard).record_ref.product_id, self.product)

    def test_disabled_grouping_autonumbers_blank_names(self):
        wizard = self.gateway([["", "GW-PRODUCT", "1"], ["", "GW-OTHER", "1"]], headers=["name", "product_id", "product_qty"], group_blank_names=False)
        wizard.action_generate_defaults()
        wizard.action_accept_import()
        self.assertEqual(set(self.results(wizard).mapped("status")), {"success"})
        self.assertEqual(len(set(self.results(wizard).mapped("identifier"))), 2)
